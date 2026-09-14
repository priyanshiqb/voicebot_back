"""Campaign reports end to end: request, build, download, sweep.

Storage is mocked at ``core.services.campaign_reports.AzureBlobStorageService``. The bytes
passed to ``upload_file`` are captured and opened with openpyxl, so these assertions are
about the file a user would actually receive, not about which functions got called.
"""

from datetime import datetime, timedelta, timezone as dt_timezone
from io import BytesIO
from unittest import mock

from django.test import override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import status
from rest_framework.test import APIClient, APITestCase

from core.models.agent import Agent
from core.models.agent_call_log import AgentCallLog
from core.models.campaign import Campaign
from core.models.campaign_member import CampaignMember
from core.models.campaign_report import CampaignReport
from core.models.contact import Contact
from core.models.in_app_notification import InAppNotification
from core.models.notification_log import NotificationLog
from core.services import campaign_reports as report_service
from core.tasks import generate_campaign_report
from tenant.models.Tenant import Tenant
from tenant.models.TenantUserMapping import TenantUserMapping
from user.models import CustomUser as User

REPORTS = "/api/v1/core/campaign-reports/"
STORAGE = "core.services.campaign_reports.AzureBlobStorageService"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

Status = CampaignReport.Status
Format = CampaignReport.Format
MemberStatus = CampaignMember.Status


class ReportTestCase(APITestCase):
    def setUp(self):
        self.tenant = Tenant.objects.create(name="Acme", subdomain="acme-reports")
        self.other_tenant = Tenant.objects.create(name="Globex", subdomain="globex-reports")

        self.user = User.objects.create_user(
            username="rita", email="rita@example.com", password="password",
            first_name="Rita", last_name="Rao",
        )
        self.outsider = User.objects.create_user(
            username="otto", email="otto@example.com", password="password"
        )
        TenantUserMapping.objects.create(user=self.user, tenant=self.tenant)
        TenantUserMapping.objects.create(user=self.outsider, tenant=self.other_tenant)

        self.api = APIClient()
        self.api.force_authenticate(user=self.user)
        self.outsider_api = APIClient()
        self.outsider_api.force_authenticate(user=self.outsider)

        self.agent = Agent.objects.create(
            tenant=self.tenant, bot_name="Sales Bot", persona_name="Persona"
        )
        self.campaign = Campaign.objects.create(
            tenant=self.tenant, agent=self.agent, name="Rabi Season",
            status=Campaign.Status.RUNNING, schedule_timezone="Asia/Kolkata",
        )

        # 03:58 UTC is 09:28 in Kolkata.
        self.called_at = datetime(2026, 9, 4, 3, 58, tzinfo=dt_timezone.utc)
        self.completed = self.member(
            "Ada", "Lovelace", "+14155552671", status=MemberStatus.COMPLETED,
            outcome="information provided", attempts=1, last_call_at=self.called_at,
        )
        self.suppressed = self.member(
            "Blocked", "Person", "+14155552672", status=MemberStatus.SUPPRESSED,
            reason=CampaignMember.SuppressionReason.DO_NOT_CALL,
        )
        self.failed = self.member(
            "Grace", "Hopper", "+14155552673", status=MemberStatus.FAILED,
            error="No answer after 3 attempts", attempts=3,
            last_call_at=self.called_at,
        )
        AgentCallLog.objects.create(
            tenant=self.tenant, agent=self.agent, provider_call_id="call-report-1",
            caller_number="+18005550100", callee_number=self.completed.phone_e164,
            direction=AgentCallLog.CallDirection.OUTBOUND, duration_seconds=95,
            credits_charged=10, campaign_member=self.completed,
        )
        self.campaign.recompute_counts()

    # --- helpers ------------------------------------------------------------

    def member(self, first, last, phone, *, status, reason="", outcome="", error="",
               attempts=0, last_call_at=None):
        contact = Contact.objects.create(
            tenant=self.tenant, first_name=first, last_name=last,
            phone_number=phone, company="Acme Traders",
        )
        return CampaignMember.objects.create(
            tenant=self.tenant, campaign=self.campaign, contact=contact,
            phone_e164=contact.phone_e164, status=status, suppression_reason=reason,
            outcome=outcome, error_message=error, attempt_count=attempts,
            last_call_at=last_call_at,
        )

    def report(self, report_format=Format.XLSX, **fields):
        return CampaignReport.objects.create(
            tenant=self.tenant, campaign=self.campaign, report_format=report_format,
            created_by=self.user, updated_by=self.user, **fields,
        )

    def build(self, report):
        """Build through the service with storage mocked. Returns (report, captured, storage)."""
        captured = {}

        def upload(key, file_obj, content_type=None):
            captured.update(key=key, data=file_obj.read(), content_type=content_type)
            return key

        with mock.patch(STORAGE) as storage_cls:
            storage = storage_cls.return_value
            storage.upload_file.side_effect = upload
            report_service.build_report(str(report.pk))

        report.refresh_from_db()
        return report, captured, storage

    def workbook(self, captured):
        return load_workbook(BytesIO(captured["data"]))

    def member_rows(self, workbook):
        rows = list(workbook["Members"].iter_rows(values_only=True))
        header = rows[0]
        return [dict(zip(header, row)) for row in rows[1:]]

    def overview(self, workbook):
        values = {}
        for row in workbook["Campaign Overview"].iter_rows(values_only=True):
            if row and row[0] is not None and len(row) > 1 and row[1] is not None:
                values[row[0]] = row[1]
        return values


# ----------------------------------------------------------------------------
# Requesting
# ----------------------------------------------------------------------------


class RequestReportTests(ReportTestCase):
    def post(self, client=None, **payload):
        url = f"/api/v1/core/campaigns/{self.campaign.pk}/reports/"
        # The patch is the *outer* context: on-commit callbacks run when the inner
        # context exits, and they must still see the patched `delay`.
        with mock.patch("core.tasks.generate_campaign_report.delay") as delay:
            with self.captureOnCommitCallbacks(execute=True):
                response = (client or self.api).post(url, payload, format="json")
        return response, delay

    def test_returns_202_with_a_pending_report_and_enqueues_once(self):
        response, delay = self.post()

        self.assertEqual(response.status_code, status.HTTP_202_ACCEPTED)
        self.assertEqual(response.data["data"]["status"], Status.PENDING)
        report = CampaignReport.objects.get()
        delay.assert_called_once_with(str(report.pk))

    def test_the_requester_is_recorded(self):
        self.post()

        self.assertEqual(CampaignReport.objects.get().created_by, self.user)

    def test_the_default_format_is_xlsx(self):
        self.post()

        self.assertEqual(CampaignReport.objects.get().report_format, Format.XLSX)

    def test_another_tenants_campaign_is_404_and_queues_nothing(self):
        response, delay = self.post(client=self.outsider_api)

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)
        self.assertFalse(CampaignReport.objects.exists())
        delay.assert_not_called()

    def test_an_unknown_format_is_400(self):
        response, _delay = self.post(report_format="CSV")

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertFalse(CampaignReport.objects.exists())

    def test_two_requests_make_two_reports(self):
        self.post()
        self.post()

        self.assertEqual(CampaignReport.objects.count(), 2)

    def test_a_pdf_can_be_requested(self):
        response, _delay = self.post(report_format=Format.PDF)

        self.assertEqual(response.status_code, status.HTTP_202_ACCEPTED)
        self.assertEqual(CampaignReport.objects.get().report_format, Format.PDF)


# ----------------------------------------------------------------------------
# Building the PDF
# ----------------------------------------------------------------------------


class BuildPdfTests(ReportTestCase):
    def test_a_pdf_report_is_a_pdf(self):
        report, captured, _storage = self.build(self.report(Format.PDF))

        self.assertEqual(report.status, Status.COMPLETED)
        self.assertTrue(report.file_name.endswith(".pdf"))
        self.assertEqual(captured["content_type"], "application/pdf")
        self.assertTrue(captured["data"].startswith(b"%PDF"))
        self.assertEqual(report.rows_included, 3)

    @override_settings(CAMPAIGN_REPORT_PDF_MAX_ROWS=2)
    def test_the_pdf_is_capped_and_the_workbook_is_not(self):
        pdf, _captured, _storage = self.build(self.report(Format.PDF))
        xlsx, _captured, _storage = self.build(self.report(Format.XLSX))

        self.assertEqual((pdf.rows_included, pdf.total_rows), (2, 3))
        self.assertEqual((xlsx.rows_included, xlsx.total_rows), (3, 3))

    def test_markup_control_characters_and_long_text_do_not_fail_the_pdf(self):
        """Paragraph text is parsed as XML, and a cell taller than a page is a LayoutError.
        Either would fail the whole document over one member."""
        contact = self.completed.contact
        contact.first_name = "<b>Ada & Co"
        contact.save()
        self.failed.error_message = "line one\nline two\x0b" + "x" * 5000
        self.failed.save()

        report, captured, _storage = self.build(self.report(Format.PDF))

        self.assertEqual(report.status, Status.COMPLETED)
        self.assertTrue(captured["data"].startswith(b"%PDF"))


# ----------------------------------------------------------------------------
# Building the workbook
# ----------------------------------------------------------------------------


class BuildWorkbookTests(ReportTestCase):
    def test_a_completed_report_records_its_file(self):
        report, captured, _storage = self.build(self.report())

        self.assertEqual(report.status, Status.COMPLETED)
        self.assertEqual(report.file_key, captured["key"])
        self.assertTrue(report.file_name.endswith(".xlsx"))
        self.assertEqual(report.file_size_bytes, len(captured["data"]))
        self.assertEqual(report.rows_included, 3)
        self.assertEqual(report.total_rows, 3)
        self.assertIsNotNone(report.expires_at)
        self.assertEqual(captured["content_type"], XLSX_TYPE)

    def test_the_workbook_has_the_overview_then_the_members(self):
        _report, captured, _storage = self.build(self.report())

        self.assertEqual(
            self.workbook(captured).sheetnames, ["Campaign Overview", "Members"]
        )

    def test_every_member_is_a_row_including_the_suppressed(self):
        _report, captured, _storage = self.build(self.report())

        rows = self.member_rows(self.workbook(captured))
        self.assertEqual(len(rows), 3)
        suppressed = next(r for r in rows if r["Name"] == "Blocked Person")
        self.assertEqual(suppressed["Suppression reason"], "On the do-not-call list")

    def test_the_phone_is_the_number_frozen_at_materialisation(self):
        """What was dialled — not what the contact has been edited to since."""
        contact = self.completed.contact
        contact.phone_number = "+14155559999"
        contact.save()

        _report, captured, _storage = self.build(self.report())

        ada = next(
            r for r in self.member_rows(self.workbook(captured)) if r["Name"] == "Ada Lovelace"
        )
        self.assertEqual(ada["Phone"], "+14155552671")

    def test_call_figures_come_from_the_call_log(self):
        _report, captured, _storage = self.build(self.report())

        rows = {r["Name"]: r for r in self.member_rows(self.workbook(captured))}
        self.assertEqual(rows["Ada Lovelace"]["Duration (seconds)"], 95)
        self.assertEqual(rows["Ada Lovelace"]["Credits charged"], 10)
        self.assertEqual(rows["Ada Lovelace"]["Provider call ID"], "call-report-1")
        # No call log, so empty cells — not zeros, not "None".
        self.assertIsNone(rows["Grace Hopper"]["Credits charged"])
        self.assertEqual(rows["Grace Hopper"]["Error"], "No answer after 3 attempts")

    def test_times_are_local_to_the_campaign_and_the_header_says_so(self):
        _report, captured, _storage = self.build(self.report())

        ada = next(
            r for r in self.member_rows(self.workbook(captured)) if r["Name"] == "Ada Lovelace"
        )
        self.assertEqual(
            ada["Last call at (Asia/Kolkata)"], datetime(2026, 9, 4, 9, 28)
        )

    def test_a_formula_in_a_contact_name_stays_inert_text(self):
        """openpyxl makes any string starting with `=` a formula; a client's Excel would
        then run it. The cell must come back as a string."""
        contact = self.completed.contact
        contact.first_name = '=HYPERLINK("http://example.com")'
        contact.save()

        _report, captured, _storage = self.build(self.report())

        sheet = self.workbook(captured)["Members"]
        header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
        name_column = header.index("Name")
        cells = [row[name_column] for row in sheet.iter_rows(min_row=2)]
        formula = next(cell for cell in cells if str(cell.value).startswith("="))
        self.assertEqual(formula.data_type, "s")

    def test_control_characters_do_not_fail_the_report(self):
        """openpyxl refuses them outright — one stray byte would fail every row."""
        self.completed.outcome = "callback\x0brequested"
        self.completed.save()

        report, captured, _storage = self.build(self.report())

        self.assertEqual(report.status, Status.COMPLETED)
        outcomes = [r["Outcome"] for r in self.member_rows(self.workbook(captured))]
        self.assertIn("callbackrequested", outcomes)

    def test_an_empty_campaign_is_a_header_only_file_not_a_failure(self):
        empty = Campaign.objects.create(
            tenant=self.tenant, agent=self.agent, name="Nobody Here",
            status=Campaign.Status.DRAFT,
        )
        report = CampaignReport.objects.create(
            tenant=self.tenant, campaign=empty, created_by=self.user, updated_by=self.user
        )

        report, captured, _storage = self.build(report)

        self.assertEqual(report.status, Status.COMPLETED)
        self.assertEqual(self.member_rows(self.workbook(captured)), [])

    def test_the_overview_reconciles_including_cancelled(self):
        self.member(
            "Carl", "Cancel", "+14155552674", status=MemberStatus.CANCELLED
        )

        _report, captured, _storage = self.build(self.report())

        overview = self.overview(self.workbook(captured))
        self.assertEqual(overview["Total members"], 4)
        self.assertEqual(overview["Cancelled"], 1)
        self.assertEqual(overview["Suppressed"], 1)
        self.assertEqual(
            overview["Completed"] + overview["Failed"] + overview["Cancelled"]
            + overview["Suppressed"],
            overview["Total members"],
        )
        self.assertEqual(overview["Credits charged"], 10)

    def test_the_requester_is_notified(self):
        self.build(self.report())

        notification = InAppNotification.objects.get(user=self.user)
        self.assertEqual(notification.notification_type, NotificationLog.NotificationType.REPORTS)
        self.assertIn("ready", notification.title)

    def test_a_report_that_is_not_pending_is_not_rebuilt(self):
        """A task delivered twice must not build the same report twice."""
        report, _captured, _storage = self.build(self.report())

        with mock.patch(STORAGE) as storage_cls:
            report_service.build_report(str(report.pk))

        storage_cls.return_value.upload_file.assert_not_called()


# ----------------------------------------------------------------------------
# Failure
# ----------------------------------------------------------------------------


class FailureTests(ReportTestCase):
    def test_a_failure_hands_the_report_back_and_leaves_no_file(self):
        """Back to PENDING so the retry can claim it; no file_key, nothing downloadable."""
        report = self.report()

        with mock.patch(STORAGE) as storage_cls:
            storage_cls.return_value.upload_file.side_effect = RuntimeError("storage down")
            with self.assertRaises(RuntimeError):
                report_service.build_report(str(report.pk))

        report.refresh_from_db()
        self.assertEqual(report.status, Status.PENDING)
        self.assertEqual(report.file_key, "")
        self.assertFalse(report.is_downloadable)

    def test_a_transient_failure_retries(self):
        report = self.report()

        with mock.patch(STORAGE) as storage_cls:
            storage_cls.return_value.upload_file.side_effect = RuntimeError("storage down")
            with mock.patch.object(
                generate_campaign_report, "retry", side_effect=RuntimeError("retried")
            ) as retry:
                with self.assertRaises(RuntimeError):
                    generate_campaign_report(str(report.pk))

        retry.assert_called_once()
        report.refresh_from_db()
        self.assertEqual(report.status, Status.PENDING)

    def test_exhausted_retries_fail_the_report_and_tell_the_requester(self):
        report = self.report()

        with mock.patch(STORAGE) as storage_cls:
            storage_cls.return_value.upload_file.side_effect = RuntimeError("storage down")
            # `.run`, not a direct call: `__call__` pushes a fresh request of its own,
            # which would reset `retries` to 0 and hide the last-attempt branch.
            generate_campaign_report.push_request(
                retries=generate_campaign_report.max_retries
            )
            try:
                with self.assertLogs("core.tasks.failures", level="ERROR"):
                    generate_campaign_report.run(str(report.pk))
            finally:
                generate_campaign_report.pop_request()

        report.refresh_from_db()
        self.assertEqual(report.status, Status.FAILED)
        self.assertEqual(report.error_message, report_service.GENERIC_FAILURE_MESSAGE)
        self.assertIn(
            "could not be generated", InAppNotification.objects.get(user=self.user).title
        )

    def test_a_deleted_campaign_fails_at_once_without_a_retry(self):
        report = self.report()
        Campaign.objects.filter(pk=self.campaign.pk).update(is_deleted=True)

        with mock.patch(STORAGE):
            with mock.patch.object(generate_campaign_report, "retry") as retry:
                generate_campaign_report(str(report.pk))

        retry.assert_not_called()
        report.refresh_from_db()
        self.assertEqual(report.status, Status.FAILED)
        self.assertIn("deleted", report.error_message)

    def test_a_notification_failure_does_not_fail_the_report(self):
        with mock.patch.object(
            InAppNotification.objects, "create", side_effect=RuntimeError("no table")
        ):
            with self.assertLogs("core.services.campaign_reports", level="ERROR"):
                report, _captured, _storage = self.build(self.report())

        self.assertEqual(report.status, Status.COMPLETED)
        self.assertTrue(report.is_downloadable)


# ----------------------------------------------------------------------------
# Polling and downloading
# ----------------------------------------------------------------------------


class DownloadTests(ReportTestCase):
    def completed_report(self):
        return self.report(
            status=Status.COMPLETED, file_key="t-x/campaign-reports/a_file.xlsx",
            file_name="campaign-rabi-season.xlsx", file_size_bytes=1234,
            completed_at=timezone.now(), expires_at=timezone.now() + timedelta(days=7),
        )

    def test_a_report_that_is_not_ready_is_409(self):
        report = self.report()

        response = self.api.get(f"{REPORTS}{report.pk}/download/")

        self.assertEqual(response.status_code, status.HTTP_409_CONFLICT)

    def test_a_completed_report_returns_a_fresh_link(self):
        report = self.completed_report()

        with mock.patch(STORAGE) as storage_cls:
            storage_cls.return_value.get_presigned_download_url.return_value = (
                "https://blob.example/file?sig=1"
            )
            response = self.api.get(f"{REPORTS}{report.pk}/download/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        data = response.data["data"]
        self.assertEqual(data["download_url"], "https://blob.example/file?sig=1")
        self.assertEqual(data["file_name"], "campaign-rabi-season.xlsx")
        storage_cls.return_value.get_presigned_download_url.assert_called_once_with(
            report.file_key, report.file_name, expiry_hours=1
        )

    def test_no_download_link_is_ever_serialised_on_the_report(self):
        report = self.completed_report()

        response = self.api.get(f"{REPORTS}{report.pk}/")

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertNotIn("download_url", response.data["data"])
        self.assertNotIn("file_key", response.data["data"])
        self.assertTrue(response.data["data"]["is_downloadable"])

    def test_another_tenants_report_is_404(self):
        report = self.completed_report()

        response = self.outsider_api.get(f"{REPORTS}{report.pk}/download/")

        self.assertEqual(response.status_code, status.HTTP_404_NOT_FOUND)

    def test_the_list_is_tenant_scoped_and_filterable(self):
        mine = self.completed_report()
        self.report()

        response = self.outsider_api.get(REPORTS)
        self.assertEqual(response.data["data"], [])

        response = self.api.get(REPORTS, {"status": Status.COMPLETED})
        ids = [row["id"] for row in response.data["data"]]
        self.assertEqual(ids, [str(mine.pk)])

    def test_reports_cannot_be_edited_or_deleted(self):
        report = self.completed_report()

        self.assertEqual(
            self.api.patch(f"{REPORTS}{report.pk}/", {"status": "FAILED"}).status_code,
            status.HTTP_405_METHOD_NOT_ALLOWED,
        )
        self.assertEqual(
            self.api.delete(f"{REPORTS}{report.pk}/").status_code,
            status.HTTP_405_METHOD_NOT_ALLOWED,
        )


# ----------------------------------------------------------------------------
# Sweeps
# ----------------------------------------------------------------------------


class SweepTests(ReportTestCase):
    def test_an_expired_report_loses_its_file_and_keeps_its_row(self):
        report = self.report(
            status=Status.COMPLETED, file_key="t-x/old.xlsx", file_name="old.xlsx",
            expires_at=timezone.now() - timedelta(minutes=1),
        )

        with mock.patch(STORAGE) as storage_cls:
            expired = report_service.expire_reports()

        storage_cls.return_value.delete_blob_direct.assert_called_once_with("t-x/old.xlsx")
        report.refresh_from_db()
        self.assertEqual(expired, 1)
        self.assertEqual(report.status, Status.EXPIRED)
        self.assertEqual(report.file_key, "")
        self.assertEqual(report.file_name, "old.xlsx")
        # It did complete; only the file is gone. Not 99, not 0.
        self.assertEqual(report.progress_percent, 100)

    def test_a_fresh_report_is_left_alone(self):
        report = self.report(
            status=Status.COMPLETED, file_key="t-x/new.xlsx",
            expires_at=timezone.now() + timedelta(days=1),
        )

        with mock.patch(STORAGE):
            report_service.expire_reports()

        report.refresh_from_db()
        self.assertEqual(report.status, Status.COMPLETED)

    def test_a_failed_file_delete_keeps_the_report_for_the_next_run(self):
        report = self.report(
            status=Status.COMPLETED, file_key="t-x/stuck.xlsx",
            expires_at=timezone.now() - timedelta(minutes=1),
        )

        with mock.patch(STORAGE) as storage_cls:
            storage_cls.return_value.delete_blob_direct.side_effect = RuntimeError("down")
            with self.assertLogs("core.services.campaign_reports", level="ERROR"):
                report_service.expire_reports()

        report.refresh_from_db()
        self.assertEqual(report.status, Status.COMPLETED)
        self.assertEqual(report.file_key, "t-x/stuck.xlsx")

    def age(self, report, **delta):
        """Backdate a report's writes. `.update()`, because `save()` would stamp
        `updated_at` with now again."""
        moment = timezone.now() - timedelta(**delta)
        CampaignReport.objects.filter(pk=report.pk).update(
            created_at=moment, updated_at=moment
        )

    def test_a_running_report_with_no_progress_for_the_window_fails(self):
        report = self.report(
            status=Status.RUNNING, started_at=timezone.now() - timedelta(hours=3)
        )
        self.age(report, hours=3)

        self.assertEqual(report_service.fail_stalled_reports(), 1)

        report.refresh_from_db()
        self.assertEqual(report.status, Status.FAILED)
        self.assertEqual(report.error_message, report_service.STALLED_MESSAGE)
        self.assertIn(
            "could not be generated", InAppNotification.objects.get(user=self.user).title
        )

    def test_a_long_report_that_is_still_progressing_is_left_alone(self):
        """Started five hours ago, but wrote progress a minute ago. Measured from the
        start, the sweep would fail a large build in the middle of writing it."""
        report = self.report(
            status=Status.RUNNING, started_at=timezone.now() - timedelta(hours=5)
        )
        self.age(report, minutes=1)

        self.assertEqual(report_service.fail_stalled_reports(), 0)
        report.refresh_from_db()
        self.assertEqual(report.status, Status.RUNNING)

    def test_a_pending_report_whose_task_never_arrived_fails(self):
        report = self.report()
        self.age(report, hours=3)

        self.assertEqual(report_service.fail_stalled_reports(), 1)
        report.refresh_from_db()
        self.assertEqual(report.status, Status.FAILED)

    def test_a_report_waiting_out_a_retry_is_left_alone(self):
        """Requested hours ago, handed back to PENDING a minute ago after a failed
        attempt. Measured from creation, the sweep would fail it in the gap before its
        last retry."""
        report = self.report()
        CampaignReport.objects.filter(pk=report.pk).update(
            created_at=timezone.now() - timedelta(hours=3),
            updated_at=timezone.now() - timedelta(minutes=1),
        )

        self.assertEqual(report_service.fail_stalled_reports(), 0)
        report.refresh_from_db()
        self.assertEqual(report.status, Status.PENDING)
