"""The campaign workbook: an overview sheet, then a streamed member sheet.

**Write-only mode, from the first line.** openpyxl's default workbook holds every cell
in memory, which for a 100,000-member campaign is how a worker gets OOM-killed.
``write_only=True`` serialises each appended row to a temp file and holds nothing. The
cost is that rows can only be appended, never revisited, and that sheet-level settings —
column widths, frozen panes — must be in place before the first row goes in.

The overview comes first because it is what a reader opens a report to see. Its figures
are computed with SQL aggregates *before* the member pass starts, so nothing here ever
needs to return to a sheet it has already written.
"""

import datetime as dt
import numbers

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from core.services.report_writers import DATETIME, NUMBER

OVERVIEW_SHEET = "Campaign Overview"
MEMBERS_SHEET = "Members"

#: A real datetime cell, sortable and groupable by month, not an ISO string.
#: The colon is escaped on purpose. Unescaped, Excel substitutes the viewer's regional
#: time separator, so the same file reads "09:02" on one machine and "09.02" on another.
DATETIME_FORMAT = "dd mmm yyyy hh\\:mm"

TITLE_FONT = Font(bold=True, size=16)
SECTION_FONT = Font(bold=True, size=12, color="1F2937")
LABEL_FONT = Font(color="4B5563")
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F2937")


def write_xlsx(content, target, *, on_progress=None, progress_every=500):
    """Write ``content`` to ``target`` (a path or a binary file). Returns rows written.

    ``on_progress(rows_written)`` is called every ``progress_every`` rows, so the caller
    can move a progress bar without this module knowing what a report row is.
    """
    workbook = Workbook(write_only=True)
    _write_overview(workbook.create_sheet(OVERVIEW_SHEET), content)
    written = _write_members(
        workbook.create_sheet(MEMBERS_SHEET), content, on_progress, progress_every
    )
    workbook.save(target)
    return written


def _text(sheet, value, *, font=None, fill=None):
    """A string cell that stays a string.

    Free text here comes from runners and from contacts, and carries two hazards:

    * **Control characters.** openpyxl refuses them with ``IllegalCharacterError``, so one
      stray ``\\x0b`` in a runner's outcome would fail the entire report, not one cell.
    * **Formulas.** openpyxl treats any string beginning with ``=`` as a formula, so a
      contact named ``=HYPERLINK(...)`` would arrive as a live link in a client's copy of
      Excel. Pinning the cell type to string keeps it inert text.
    """
    cell = WriteOnlyCell(sheet, value=ILLEGAL_CHARACTERS_RE.sub("", str(value)))
    cell.data_type = "s"
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    return cell


def _datetime(sheet, value):
    cell = WriteOnlyCell(sheet, value=value)
    cell.number_format = DATETIME_FORMAT
    return cell


def _is_number(value):
    # bool is an int in Python; a True written as 1 is a quietly wrong spreadsheet.
    return isinstance(value, numbers.Number) and not isinstance(value, bool)


def _overview_value(sheet, value):
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return _text(sheet, "Yes" if value else "No")
    if _is_number(value):
        return value
    if isinstance(value, dt.datetime):
        return _datetime(sheet, value)
    return _text(sheet, value)


def _member_value(sheet, column, value):
    # Empty means an empty cell, never "None" or "N/A" — those are strings a SUM trips
    # over and someone has to filter out by hand.
    if value is None or value == "":
        return None
    if column.kind == NUMBER and _is_number(value):
        return value
    if column.kind == DATETIME and isinstance(value, dt.datetime):
        return _datetime(sheet, value)
    return _text(sheet, value)


def _write_overview(sheet, content):
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 52

    sheet.append([_text(sheet, content.title, font=TITLE_FONT)])
    sheet.append([_text(sheet, content.subtitle, font=LABEL_FONT)])

    for section, pairs in content.overview:
        sheet.append([])
        sheet.append([_text(sheet, section, font=SECTION_FONT)])
        for label, value in pairs:
            sheet.append([_text(sheet, label, font=LABEL_FONT), _overview_value(sheet, value)])


def _write_members(sheet, content, on_progress, progress_every):
    # Both before the first row: write-only mode emits the sheet's head on the first
    # append and cannot go back for it.
    for index, column in enumerate(content.columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column.width
    sheet.freeze_panes = "A2"

    sheet.append([
        _text(sheet, column.header, font=HEADER_FONT, fill=HEADER_FILL)
        for column in content.columns
    ])

    written = 0
    for row in content.rows:
        sheet.append([
            _member_value(sheet, column, value)
            for column, value in zip(content.columns, row)
        ])
        written += 1
        if on_progress is not None and written % progress_every == 0:
            on_progress(written)

    # The autofilter is written at the tail of the sheet, so it can be sized now that the
    # row count is known. A 20,000-row sheet is then usable the moment it opens.
    last_column = get_column_letter(len(content.columns))
    sheet.auto_filter.ref = f"A1:{last_column}{written + 1}"
    return written
